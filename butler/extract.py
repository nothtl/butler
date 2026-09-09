"""Text extraction (feature 13).

Pulls searchable text out of common file types. Extracted text powers both
FTS5 content search (features 4/5/9) and semantic embeddings (feature 14).
"""

from __future__ import annotations

import csv
import html.parser
import json
import os
import re
import shutil
import subprocess
from typing import Any

# Image types that only yield text via OCR (feature: scanned docs / photos).
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".gif"}

TEXT_EXT = {
    ".txt", ".md", ".rst", ".log", ".csv", ".tsv", ".json", ".xml", ".yaml",
    ".yml", ".toml", ".ini", ".cfg", ".conf", ".html", ".htm", ".py", ".js",
    ".ts", ".c", ".cpp", ".h", ".java", ".rs", ".go", ".rb", ".php", ".sh",
    ".sql", ".css", ".tex",
}


class _StripHTML(html.parser.HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        s = data.strip()
        if s:
            self.parts.append(s)


def _clean(text: str | None) -> str:
    if not text:
        return ""
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t\u200b]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(path: str, ocr: bool = False) -> dict[str, Any]:
    """Return {'text': str, 'meta': dict} for a file path.

    Never raises for unsupported/corrupt content — returns empty text.
    When ``ocr`` is True and a tesseract binary is present, scanned PDFs and
    images are run through OCR; otherwise text extraction degrades gracefully.
    """
    ext = os.path.splitext(path)[1].lower()
    base = {"text": "", "meta": {}}
    if ext in TEXT_EXT:
        return base | {"text": _clean(_read_text_file(path))}
    if ext == ".pdf":
        res = base | _pdf(path)
        if ocr and not res.get("text"):
            return res | {"text": _clean(_pdf_ocr(path)), "meta": res.get("meta", {})}
        return res
    if ext in IMAGE_EXT and ocr:
        return base | {"text": _clean(_image_ocr(path))}
    if ext in (".docx",):
        return base | _docx(path)
    if ext == ".pptx":
        return base | _pptx(path)
    if ext in (".epub", ".mobi"):
        return base | {"text": _clean(_epub(path))}
    return base


def _read_text_file(path: str) -> str:
    raw = open(path, "rb").read()
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _pdf(path: str) -> dict[str, Any]:
    try:
        import pymupdf  # imported submodule (new API)
    except ImportError:
        try:
            import fitz  # type: ignore
            pymupdf = fitz
        except ImportError:
            return {"text": "", "meta": {}}
    try:
        doc = pymupdf.open(path)
        parts: list[str] = []
        for page in doc:
            parts.append(page.get_text("text"))
        meta = {k: doc.metadata.get(k, "") for k in ("title", "author", "subject")}
        return {"text": _clean("\n".join(parts)), "meta": meta}
    except Exception:
        return {"text": "", "meta": {}}


def _docx(path: str) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError:
        return {"text": "", "meta": {}}
    try:
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        # include table cells
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        meta = {}
        cp = d.core_properties
        meta["title"] = cp.title or ""
        meta["author"] = cp.author or ""
        return {"text": _clean("\n".join(parts)), "meta": meta}
    except Exception:
        return {"text": "", "meta": {}}


def _pptx(path: str) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError:
        return {"text": "", "meta": {}}
    try:
        prs = Presentation(path)
        parts: list[str] = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        parts.append("".join(run.text for run in para.runs))
        return {"text": _clean("\n".join(parts)), "meta": {}}
    except Exception:
        return {"text": "", "meta": {}}


def _epub(path: str) -> str:
    try:
        import zipfile
        from xml.etree import ElementTree as ET

        with zipfile.ZipFile(path) as zf:
            texts = []
            for name in zf.namelist():
                if name.endswith((".xhtml", ".html", ".htm")):
                    raw = zf.read(name).decode("utf-8", errors="replace")
                    parser = _StripHTML()
                    parser.feed(raw)
                    texts.append(" ".join(parser.parts))
            return "\n".join(texts)
    except Exception:
        return ""


def _ocr_ready() -> bool:
    if not shutil.which("tesseract"):
        return False
    try:
        import pytesseract  # noqa: F401
        return True
    except ImportError:
        return False


def _image_ocr(path: str) -> str:
    if not _ocr_ready():
        return ""
    try:
        import pytesseract
        from PIL import Image
        with Image.open(path) as im:
            return pytesseract.image_to_string(im)
    except Exception:
        return ""


def _pdf_ocr(path: str) -> str:
    """Rasterise each PDF page and OCR it (for scanned, textless PDFs)."""
    if not _ocr_ready():
        return ""
    try:
        import pymupdf
    except ImportError:  # pragma: no cover
        try:
            import fitz  # type: ignore
            pymupdf = fitz
        except ImportError:
            return ""
    try:
        import pytesseract
        from PIL import Image
        import io
        doc = pymupdf.open(path)
        parts: list[str] = []
        for page in doc:
            pix = page.get_pixmap(dpi=200)
            with Image.open(io.BytesIO(pix.tobytes("png"))) as im:
                parts.append(pytesseract.image_to_string(im))
        return "\n".join(parts)
    except Exception:
        return ""


def chunk_text(text: str, size: int = 800, overlap: int = 100) -> list[str]:
    """Split text into overlapping chunks suitable for embedding/FTS.

    Chunking is paragraph-and-sentence aware (lightweight, no NLP deps).
    """
    text = text.strip()
    if not text:
        return []
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    sentences: list[str] = []
    for para in paragraphs:
        sentences.extend(s.strip() for s in re.split(r"(?<=[.!?])\s+", para) if s.strip())

    chunks: list[str] = []
    cur = ""
    for sentence in sentences:
        if len(cur) + len(sentence) + 1 > size and cur:
            chunks.append(cur)
            cur = cur[-overlap:]
        cur = (cur + " " + sentence).strip()
    if cur:
        chunks.append(cur)
    return chunks
